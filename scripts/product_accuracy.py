import re
import random
import sqlite3
from collections import defaultdict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

DB_PATH = "my_database.db"
OUT_PATH = "appendix_results.docx"
STORES = ["Slot", "Jumia", "Konga"]   # must match the names in your stores table
PER_CATEGORY = 10                     # products per category
MIN_STORES = 2                        # products must be listed in at least this many stores
SEED = 42                             # same seed = same products every run

QUERY = """
SELECT p.id, c.name, b.name, p.model, p.ram, p.storage, s.name, sp.name
FROM store_products sp
JOIN products   p ON p.id = sp.product_id
JOIN stores     s ON s.id = sp.store_id
JOIN brands     b ON b.id = p.brand_id
JOIN categories c ON c.id = p.category_id
WHERE p.is_deleted = 0
"""

# ---------------------------------------------------------------- reading a listing name
def tok(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()

def contains(text, phrase):
    return bool(tok(phrase)) and f" {tok(phrase)} " in f" {tok(text)} "

def flat(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())

SIZE = re.compile(r"(\d+(?:\.\d+)?)\s*(gb|tb)\b", re.I)

def read_sizes(text):
    """Return (ram, rom) as the listing name states them, or None if not stated."""
    found = []
    for m in SIZE.finditer(text):
        num, unit = float(m.group(1)), m.group(2).upper()
        after = text[m.end(): m.end() + 5].lower()
        before = text[max(0, m.start() - 5): m.start()].lower()
        label = None
        for chunk in (after, before):
            if "ram" in chunk:
                label = "ram"
            elif any(k in chunk for k in ("rom", "ssd", "storage")):
                label = "rom"
            if label:
                break
        shown = f"{int(num) if num.is_integer() else num}{unit}"
        found.append((num * (1024 if unit == "TB" else 1), shown, label))

    ram = next((s for _, s, l in found if l == "ram"), None)
    rom = next((s for _, s, l in found if l == "rom"), None)
    rest = sorted((f for f in found if f[2] is None and f[1] not in (ram, rom)), key=lambda f: f[0])

    if ram is None and rom is None:
        if len(rest) >= 2:
            ram, rom = rest[0][1], rest[-1][1]
        elif len(rest) == 1:
            if rest[0][0] >= 32:
                rom = rest[0][1]
            else:
                ram = rest[0][1]
    elif ram is None and rest:
        ram = rest[0][1]
    elif rom is None and rest:
        rom = rest[-1][1]
    return ram, rom

def evaluate(canon, listing_name):
    """Compare the system product with what the listing name says."""
    brand, model, ram, rom = canon
    l_ram, l_rom = read_sizes(listing_name)
    shown, res = {}, {}

    if brand and contains(listing_name, brand):
        shown["brand"], res["brand"] = brand, True
    else:
        shown["brand"], res["brand"] = "—", None          # not stated in the listing

    if model and contains(listing_name, model):
        shown["model"], res["model"] = model, True
    else:
        shown["model"], res["model"] = "not found", False

    for key, system_value, listing_value in (("ram", ram, l_ram), ("rom", rom, l_rom)):
        if listing_value is None:
            shown[key], res[key] = "—", None               # not stated: excluded from the score
        else:
            shown[key] = listing_value
            res[key] = flat(system_value) == flat(listing_value)

    scored = [res[k] for k in ("model", "ram", "rom") if res[k] is not None]
    return shown, res, sum(scored) / len(scored)

# ---------------------------------------------------------------- pick the products
conn = sqlite3.connect(DB_PATH)
rows = conn.execute(QUERY).fetchall()
conn.close()

products = {}
for pid, cat, brand, model, ram, rom, store, lname in rows:
    store = next((s for s in STORES if s.lower() == str(store).strip().lower()), None)
    if not store:
        continue
    e = products.setdefault(pid, {"cat": cat, "canon": (brand, model, ram, rom), "listings": {}})
    e["listings"].setdefault(store, lname)

random.seed(SEED)
candidates = defaultdict(list)
for pid, e in products.items():
    if len(e["listings"]) >= MIN_STORES:
        candidates[e["cat"]].append(e)

selected = {}
for cat, items in candidates.items():
    random.shuffle(items)
    items.sort(key=lambda e: -len(e["listings"]))          # products in all 3 stores come first
    picked = items[:PER_CATEGORY]
    picked.sort(key=lambda e: tuple(str(v or "") for v in e["canon"]))
    selected[cat] = picked
    if len(picked) < PER_CATEGORY:
        print(f"Only {len(picked)} comparable products in '{cat}'")

# ---------------------------------------------------------------- score everything first
def blank():
    return {"n": 0, "score": 0.0, "model": [0, 0], "ram": [0, 0], "rom": [0, 0]}

def add_stats(st, res, score):
    st["n"] += 1
    st["score"] += score
    for k in ("model", "ram", "rom"):
        if res[k] is not None:
            st[k][1] += 1
            st[k][0] += bool(res[k])

def pct(h):
    return f"{100 * h[0] / h[1]:.1f}%" if h[1] else "n/a"

by_cat, by_store, overall = defaultdict(blank), defaultdict(blank), blank()
evaluated = {}
for cat, items in selected.items():
    evaluated[cat] = []
    for e in items:
        evs = []
        for store in STORES:
            if store in e["listings"]:
                shown, res, score = evaluate(e["canon"], e["listings"][store])
                evs.append((store, e["listings"][store], shown, res, score))
                for st in (by_cat[cat], by_store[store], overall):
                    add_stats(st, res, score)
        evaluated[cat].append((e, evs))

# ---------------------------------------------------------------- Word helpers
CENTER = WD_ALIGN_PARAGRAPH.CENTER

def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)

def write(cell, text, bold=False, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    run = p.add_run(str(text))
    run.font.size = Pt(9)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def flag_row(row, tag):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement(tag)
    el.set(qn("w:val"), "true")
    trPr.append(el)

def make_table(doc, headers, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, h, w in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(w)
        write(cell, h, bold=True, color="FFFFFF", align=CENTER)
        shade(cell, "1F3864")
    flag_row(t.rows[0], "w:tblHeader")                     # header repeats on every page
    return t

def add_row(t, widths):
    row = t.add_row()
    flag_row(row, "w:cantSplit")
    for cell, w in zip(row.cells, widths):
        cell.width = Inches(w)
    return row

def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True
    p.add_run(text).bold = True

def summary_table(doc, cap, first_header, groups):
    caption(doc, cap)
    widths = [3.6, 1.3, 1.3, 1.3, 1.3, 1.3]
    t = make_table(doc, [first_header, "Listings", "Model", "RAM", "ROM", "Mean score"], widths)
    for name, st in groups:
        row = add_row(t, widths)
        vals = [name, st["n"], pct(st["model"]), pct(st["ram"]), pct(st["rom"]),
                f"{st['score'] / st['n']:.2f}" if st["n"] else "n/a"]
        for i, (c, v) in enumerate(zip(row.cells, vals)):
            write(c, v, bold=(name == "Overall"), align=None if i == 0 else CENTER)
            if name == "Overall":
                shade(c, "DDEBF7")

# ---------------------------------------------------------------- build the document
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)      # A4 landscape
sec.left_margin = sec.right_margin = Cm(2)
sec.top_margin = sec.bottom_margin = Cm(2)

doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(10)
for name in ("Heading 1", "Heading 2"):
    doc.styles[name].font.name = "Times New Roman"
    doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)

doc.add_heading("Appendix: Full Evaluation Results", level=1)
doc.add_paragraph(
    f"For each product category, {PER_CATEGORY} products listed in at least {MIN_STORES} of the "
    f"{len(STORES)} stores were compared. The shaded row is the system's canonical product. The rows "
    "beneath it show what each store's listing states. A dash means the listing does not state that "
    "field, and it is left out of scoring. Red marks a mismatch with the system product. A listing's "
    "score is the share of its stated model, RAM and ROM values that match the system product."
)

summary_table(doc, "Table A1: Accuracy by category", "Category",
              [(c, by_cat[c]) for c in sorted(by_cat)] + [("Overall", overall)])
summary_table(doc, "Table A2: Accuracy by store", "Store",
              [(s, by_store[s]) for s in STORES if s in by_store] + [("Overall", overall)])

HEADERS = ["Ref", "Source", "Product / listing name", "Brand", "Model", "RAM", "ROM", "Score"]
WIDTHS = [0.55, 0.8, 4.0, 1.1, 1.5, 0.7, 0.8, 0.65]
RED = "C00000"

for cat in sorted(evaluated):
    doc.add_page_break()
    doc.add_heading(cat, level=2)
    caption(doc, f"Table A3.{sorted(evaluated).index(cat) + 1}: {cat} ({len(evaluated[cat])} products)")
    t = make_table(doc, HEADERS, WIDTHS)

    for n, (e, evs) in enumerate(evaluated[cat], start=1):
        brand, model, ram, rom = e["canon"]
        name = " ".join(str(x) for x in (brand, model, ram, rom) if x)
        mean = sum(ev[4] for ev in evs) / len(evs)
        block = []

        head = add_row(t, WIDTHS)
        block.append(head)
        vals = [f"P{n:02d}", "System", name, brand or "—", model or "—", ram or "—", rom or "—", f"{mean:.2f}"]
        for i, (c, v) in enumerate(zip(head.cells, vals)):
            write(c, v, bold=True, align=None if i in (1, 2, 3, 4) else CENTER)
            shade(c, "DDEBF7")

        for store, lname, shown, res, score in evs:
            row = add_row(t, WIDTHS)
            block.append(row)
            cells = row.cells
            write(cells[0], "")
            write(cells[1], store)
            write(cells[2], lname)
            for idx, key in ((3, "brand"), (4, "model"), (5, "ram"), (6, "rom")):
                bad = res[key] is False
                write(cells[idx], shown[key], color=RED if bad else None, bold=bad,
                      align=None if idx in (3, 4) else CENTER)
            write(cells[7], f"{score:.2f}", align=CENTER)

        for row in block[:-1]:                              # keep each product's block on one page
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True

doc.save(OUT_PATH)
total = sum(len(v) for v in evaluated.values())
print(f"Saved {OUT_PATH}: {total} products across {len(evaluated)} categories")